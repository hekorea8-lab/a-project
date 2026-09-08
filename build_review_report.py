from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "outputs" / "회계세무챗봇_전문개발_개선진단보고서.docx"


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    props.append(element)


def borders(table, color="D9D9D9", size="6"):
    tbl = table._tbl
    props = tbl.tblPr
    border = props.first_child_found_in("w:tblBorders")
    if border is None:
        border = OxmlElement("w:tblBorders")
        props.append(border)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        node = border.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            border.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_cell_text(cell, text, bold=False, color="000000", size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = "맑은 고딕"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = True
    borders(table)
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True, color="FFFFFF", size=8.5)
        shade(table.rows[0].cells[i], "1F4E79")
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, size=8.5)
            if row_index % 2 == 1:
                shade(cells[i], "F4F8FB")
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.font.name = "맑은 고딕"
    return p


def add_paragraph(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.8)
section.bottom_margin = Cm(1.6)
section.left_margin = Cm(2.0)
section.right_margin = Cm(2.0)

styles = doc.styles
styles["Normal"].font.name = "맑은 고딕"
styles["Normal"].font.size = Pt(10)
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
for style_name, size in (("Title", 24), ("Heading 1", 16), ("Heading 2", 12), ("Heading 3", 10.5)):
    style = styles[style_name]
    style.font.name = "맑은 고딕"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("회계·세무 지식 챗봇 전문 개발 진단 보고서 | 2026-09-08").font.size = Pt(8)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.add_run("회계 세무 지식 챗봇 전문 개발 진단 보고서")
sub = doc.add_paragraph()
sub.paragraph_format.space_after = Pt(18)
sub.add_run("현재 구조 기준 부족한 부분과 개선 우선순위").font.size = Pt(12)
sub.runs[0].font.color.rgb = RGBColor(89, 89, 89)

add_paragraph(doc, "이 보고서는 현재 프로젝트의 코드를 수정하지 않고, FastAPI 앱 구조·RAG 검색·임베딩·LLM 답변·계산 기능·보안·테스트·운영성 관점에서 개선 필요사항을 점검한 결과입니다. 결론부터 말하면, 핵심 기능은 이미 상당 부분 구현되어 있으나 운영 서비스로 안정화하려면 벡터 검색의 실제 답변 반영, 모놀리식 구조 분리, 인증과 업로드 보안, 주장 단위 검증, 독립 테스트 체계가 우선 보완되어야 합니다.")
add_paragraph(doc, "검토 기준: 현재 소스(app.py, backend/main.py), PRD, 실제 실행 상태, 지식DB·임베딩 상태, 내장 품질 테스트 결과를 확인했습니다. 현 시점에서 확인된 실행 상태는 API 정상, SQLite 지식 청크 28,881개, PostgreSQL pgvector 임베딩 45,049개, 대상 문서 411개, 내장 품질 테스트 26개 통과입니다.")

add_heading(doc, "1 현재 시스템 진단 요약", 1)
add_table(doc, ["영역", "현재 확인 상태", "진단"], [
    ("RAG 검색", "질문 분석, rewrite, FTS5 BM25, 구조화 검색, 선택적 벡터 검색, rerank, grounding 구현", "기능 기반은 양호하나 검색 경로가 복잡하고 평가 지표가 부족함"),
    ("임베딩", "45,049개 저장, text-embedding-3-large, 3,072차원, HNSW", "현재 mode가 shadow라 최종 답변에 벡터 결과가 적극 반영되지 않음"),
    ("LLM", "단순 조회·검색 계획·전문가 답변 프롬프트 분리", "HumanMessage 본문 중심이며 주장 단위 사실 검증은 제한적"),
    ("UI", "기존 페이지·챗봇·로딩·근거 표시 유지", "기능은 충분하나 진행률·예상시간이 실제 단계 기반이 아닌 시간 추정 중심"),
    ("계산", "기존 세무 계산과 회계 계산 스킬 카탈로그 확장", "챗봇 계산과 별도 세액 계산 API의 기능 범위가 분리되어 있음"),
    ("운영", "health, refresh status, embedding status 제공", "인증·권한·rate limit·관측성·자동화가 운영 수준으로 부족"),
    ("테스트", "app.py 내 unittest 26개", "독립 테스트 파일과 회귀용 고정 데이터셋이 부족"),
], [3.0, 6.3, 7.2])

add_heading(doc, "2 가장 영향도가 큰 문제 TOP 10", 1)
add_table(doc, ["순위", "문제", "영향도", "권장 조치"], [
    ("1", "임베딩 검색이 shadow mode로 운영되어 최종 답변 근거에 사실상 반영되지 않음", "매우 높음", "검증셋으로 hybrid 전환 여부를 비교하고, 안전한 문서 유형부터 단계적으로 활성화"),
    ("2", "app.py가 약 8,116줄·259개 정의를 포함하는 단일 파일이며 backend/main.py에도 유사 구현이 존재", "매우 높음", "즉시 전면 리팩토링하지 말고 검색·답변·API·UI를 단계별 모듈로 분리"),
    ("3", "인증·권한·업로드 정책이 코드상 명확히 확인되지 않음", "매우 높음", "관리자·갱신·분석·첨부 엔드포인트에 인증과 역할별 권한 적용"),
    ("4", "LLM 출력 검증이 evidence ID와 형식 중심이며 숫자·법령 주장·계산식의 주장 단위 검증이 약함", "높음", "claim-evidence 검증과 숫자·날짜·세율 일치 검증 추가"),
    ("5", "독립된 테스트 파일이 없고 app.py 실행형 self-check에 테스트가 집중", "높음", "tests/ 아래 pytest 또는 unittest 파일과 고정 RAG 평가셋 추가"),
    ("6", "수동 갱신·임베딩 색인·승인 버전 관리가 운영 프로세스로 정리되지 않음", "높음", "수집-정제-승인-색인-롤백 상태와 스케줄링 정의"),
    ("7", "문서·청크·임베딩·메타데이터의 동기화 무결성 검사가 제한적", "높음", "content_hash 기반 불일치 리포트와 삭제·구버전 정리 정책 추가"),
    ("8", "검색 로그는 있으나 BM25·구조화 검색·벡터 결과가 동일한 형태로 추적되지 않음", "중간", "검색 단계별 원본 후보와 점수, 탈락 사유, 최종 선택 사유를 표준화"),
    ("9", "모델 호출·예외 처리·응답 지연에 대한 운영 메트릭이 부족", "중간", "request_id별 latency, token, 오류, fallback, 검색 품질을 구조화 기록"),
    ("10", "계산 스킬 카탈로그와 별도 /tax-calculations API의 지원 범위가 다름", "중간", "공통 계산 엔진과 계산 스킬 registry를 단일 기준으로 확장"),
], [1.0, 7.2, 2.0, 6.3])

add_heading(doc, "3 상세 검토 결과", 1)

add_heading(doc, "3.1 RAG와 임베딩", 2)
add_paragraph(doc, "현재 구조는 SQLite document_chunks를 기준으로 FTS5 BM25·구조화 키워드 검색을 수행하고, PostgreSQL pgvector에는 45,049개 임베딩을 저장합니다. 그러나 EMBEDDING_RETRIEVAL_MODE 기본값이 shadow이며, search_hybrid_documents에서 semantic 결과를 최종 fuse 대상에서 제외하는 조건이 있습니다. 따라서 임베딩 저장 상태와 실제 답변 기여 상태가 다릅니다.")
add_paragraph(doc, "개선 방향은 곧바로 벡터 검색을 전면 활성화하는 것이 아니라, 동일한 평가셋으로 BM25 단독·벡터 단독·BM25+벡터를 비교한 뒤 분야별로 단계적 전환하는 것입니다. 법령 조문번호·세목처럼 정확 일치가 중요한 질문은 BM25 가중치를 높이고, 표현이 다양한 회계 질문은 벡터 후보를 보강 신호로 사용하는 방식이 적합합니다.")

add_heading(doc, "3.2 구조와 유지보수", 2)
add_paragraph(doc, "app.py 하나에 데이터 모델, API, HTML, JavaScript, 수집, PDF 처리, 검색, 임베딩, 그래프, LLM, 계산, 테스트가 함께 들어 있습니다. 현재 요구사항을 빠르게 반영하기에는 유리하지만, 작은 변경이 다른 흐름을 깨뜨릴 가능성이 커지고 코드 탐색·리뷰·배포 단위가 어려워집니다. backend/main.py에도 별도 구현이 있어 두 진입점의 동작 차이가 발생할 수 있습니다.")
add_paragraph(doc, "사용자 요청대로 기존 UI와 정상 기능을 유지해야 하므로 1차 조치는 삭제가 아닌 경계 설정입니다. 예를 들어 rag_pipeline.py, answer_service.py, calculation_engine.py, auth.py를 새로 두고 app.py는 호환 호출만 유지하는 점진적 추출이 안전합니다.")

add_heading(doc, "3.3 보안과 개인정보", 2)
add_paragraph(doc, "코드상 API 키는 .env에서 읽는 원칙이 지켜지고 있습니다. 다만 관리자 페이지, 지식 갱신, CSV 원장, 첨부파일, 분석 저장 API에 대한 인증·역할 검사가 명확히 보이지 않습니다. 실제 내부망 서비스라도 링크를 아는 사용자의 임의 실행, 대용량 업로드, 민감 원문 노출, 관리자 기능 오용을 방지하는 통제가 필요합니다.")
add_bullets(doc, [
    "최소 권한: 일반 사용자, 검토 담당자, 지식 관리자, 시스템 관리자 역할 분리",
    "업로드: 확장자·MIME·파일 크기·페이지 수·압축폭탄·OCR 처리시간 제한",
    "로그: 원문·첨부 텍스트·API 키를 남기지 않고 질문 해시·문서 ID·점수만 보존",
    "운영: CORS, CSRF, rate limit, 관리자 작업 감사로그, 보존기간과 삭제정책 정의",
])

add_heading(doc, "3.4 답변 품질과 검증", 2)
add_paragraph(doc, "현재는 DIRECT/PARTIAL/IRRELEVANT와 evidence_ids 검증이 구현되어 있어 무관 문서 제거의 기반은 좋습니다. 그러나 LLM이 작성한 세율·기한·금액·계산식이 실제 원문과 일치하는지, 문장별 주장이 어떤 근거의 어느 구간에서 나왔는지까지는 일반화된 검증기가 아닙니다. 특히 세율, 기업 규모별 공제율, 적용연도, 신고기한처럼 숫자 하나가 결론을 바꾸는 질문은 별도 검증이 필요합니다.")
add_paragraph(doc, "권장 검증 순서는 검색 적합성 확인 → 근거 문단 선택 → 주장 추출 → 숫자·날짜·조문 대조 → 답변 생성 → 생성 후 재검증입니다. 검증 실패 시 전체 답변을 막기보다 확인된 항목만 보여주는 부분 답변 정책을 명확히 해야 합니다.")

add_heading(doc, "3.5 계산 기능", 2)
add_paragraph(doc, "현재 회계학개론 수준의 일부 계산식과 세무 계산 스킬 카탈로그가 확장되어 있습니다. 손상차손, 처분손익, 매출총이익, 이익률, 감가상각 및 기존 세무 계산이 테스트되고 있습니다. 다만 계산 결과는 입력 순서와 가정에 의존하므로 계산 전 입력값 의미를 구조화하고, 단위·기간·세전/세후·부가세 포함 여부를 확인하는 공통 입력 검증이 필요합니다.")
add_paragraph(doc, "또한 화면의 /tax-calculations API는 현재 국가전략기술 공제, 무신고가산세, 납부지연가산세 중심이고, 자연어 챗봇 내부 계산 스킬과 별도입니다. 앞으로는 계산식을 하나의 registry에서 관리하고, 검색 근거가 필요한 계산과 순수 산술 계산을 구분해야 합니다.")

add_heading(doc, "4 승인 후 개선 작업 리스트", 1)
add_paragraph(doc, "아래 목록은 기존 UI와 정상 동작을 보존하면서 단계적으로 적용할 수 있도록 작성했습니다. 먼저 1단계만 승인해도 검색 품질과 운영 안전성의 핵심 위험을 줄일 수 있습니다.")
add_table(doc, ["단계", "작업", "대상 파일", "완료 기준"], [
    ("1A", "임베딩 shadow/hybrid 비교 평가 및 단계적 활성화", "app.py, tests/평가셋", "Recall@5·MRR·Precision@5를 BM25와 비교하고 분야별 모드 결정"),
    ("1B", "인증·역할·업로드 제한·관리자 감사로그", "app.py 또는 신규 auth 모듈", "비인가 호출 차단, 업로드 한도와 감사 기록 확인"),
    ("1C", "주장 단위 근거·숫자·날짜 검증", "app.py 또는 신규 validation 모듈", "세율·기한·금액 불일치 답변 자동 보류"),
    ("1D", "독립 테스트와 RAG 고정 평가셋", "tests/", "세무·회계 20개 이상 질문, Top 5와 지표 자동 리포트"),
    ("2A", "검색·답변·계산 경계 분리", "app.py, backend/main.py", "호환 API 유지, 기능별 단위 테스트 가능"),
    ("2B", "지식기반 버전·승인·롤백 관리", "app.py, data 스키마", "현행/구버전 구분, 재색인 실패 시 이전 버전 복귀"),
    ("2C", "관측성 강화", "app.py, 운영 설정", "request_id별 단계 latency, fallback, 검색점수, 오류 확인"),
    ("3A", "계산 스킬 registry 확장", "app.py 또는 calculation_engine.py", "회계·세무 계산식 추가 시 공통 입력·단위·근거 검증"),
], [1.2, 6.0, 4.2, 5.0])

add_heading(doc, "5 현재 강점", 1)
add_bullets(doc, [
    "사용자 질문을 구조화하고 검색용 표현으로 다시 쓰는 단계가 존재합니다.",
    "SQLite FTS5 BM25와 구조화 키워드 검색이 실제 검색 경로에 연결되어 있습니다.",
    "법령명·세목·조문·과세대상별 검색과 재산세 유형 분해 로직이 있습니다.",
    "직접 근거와 부분 근거를 구분하고 무관 문서를 제거하는 grounding 단계가 있습니다.",
    "법령·회계기준·판례·첨부자료를 서로 다른 근거 유형으로 관리합니다.",
    "임베딩 45,049개와 HNSW 인덱스가 준비되어 있어 hybrid 검색 실험 기반이 있습니다.",
    "단순 질문은 LLM 호출을 생략해 응답시간을 줄이는 경로가 있습니다.",
    "기존 기능을 보존하면서 계산 스킬과 피드백·검색 로그를 확장할 수 있는 구조가 있습니다.",
])

add_heading(doc, "6 확인된 테스트와 운영 상태", 1)
add_table(doc, ["점검 항목", "결과", "의미"], [
    ("Python 문법 검사", "통과", "현재 app.py가 실행 가능한 상태"),
    ("내장 품질 테스트", "26개 통과", "기존 시나리오와 최근 계산 기능의 기본 회귀 통과"),
    ("HTTP health", "connected", "API와 SQLite 데이터 연결 확인"),
    ("임베딩 probe", "ready", "PostgreSQL pgvector 테이블과 45,049행 확인"),
    ("SQLite knowledge DB", "28,881 chunks / 411 documents", "검색 원문과 청크가 존재"),
], [4.5, 3.5, 8.4])
add_paragraph(doc, "테스트가 통과했다는 사실은 현재 시나리오의 회귀가 없다는 뜻이지, 실제 법령 질문 전반의 검색 정확도나 운영 보안을 보장하지는 않습니다. 다음 개선의 출발점은 고정된 정답 근거를 가진 평가셋과 비인가·대용량·오류 복구 테스트입니다.")

add_heading(doc, "7 최종 의견", 1)
add_paragraph(doc, "현재 시스템은 단순 PoC를 넘어 회계·세무 검토 서비스로 확장할 수 있는 핵심 구성요소를 갖추고 있습니다. 다만 현재의 가장 큰 위험은 기능 부재보다 ‘구현된 기능이 실제 최종 경로에서 어느 정도 사용되는지’, ‘근거와 숫자가 답변 문장 단위로 검증되는지’, ‘내부 서비스로 운영될 때 접근과 원문이 보호되는지’를 한눈에 확인하기 어렵다는 점입니다.")
add_paragraph(doc, "승인 우선순위는 1) 임베딩 hybrid 평가와 RAG 평가셋, 2) 인증·업로드 보안, 3) 주장 단위 검증, 4) 독립 테스트 체계, 5) 점진적 모듈 분리 순서가 적절합니다. 이 순서는 UI를 바꾸거나 정상 기능을 삭제하지 않고도 품질과 운영 안정성을 가장 빠르게 높일 수 있습니다.")

add_heading(doc, "8 승인 요청 리스트", 1)
add_paragraph(doc, "보고서 검토 후 아래 항목의 진행 여부를 선택하면, 선택된 범위만 기존 구조를 보존하면서 구현할 수 있습니다.")
add_table(doc, ["선택", "개선 항목", "권장"], [
    ("□", "RAG 평가셋·BM25/벡터/hybrid 비교", "최우선"),
    ("□", "임베딩 shadow에서 hybrid로 단계적 전환", "평가 후 진행"),
    ("□", "인증·권한·업로드 보안", "운영 전 필수"),
    ("□", "근거·숫자·날짜 주장 단위 검증", "운영 전 필수"),
    ("□", "독립 테스트 파일과 회귀 자동화", "최우선"),
    ("□", "모듈 분리와 backend/main.py 정리", "중기"),
    ("□", "지식기반 버전·승인·롤백 관리", "중기"),
    ("□", "계산 스킬 registry 확대", "확장"),
], [1.0, 10.0, 4.0])

doc.core_properties.title = "회계 세무 지식 챗봇 전문 개발 진단 보고서"
doc.core_properties.subject = "현재 구조 기준 개선 필요사항과 우선순위"
doc.core_properties.author = "Codex"
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
